"""
Main Pipeline Orchestrator

Orchestrates the complete workflow:
PDF → OCR → Layout Analysis → Structured JSON → Template Learning
→ Section-wise Retrieval → Constrained Generation → Final Reconstruction

Extended with layout capture:
  Option 1 - PNG visual overlays per page
  Option 2 - pdfplumber table extraction to JSON
  Option 3 - Full spatial layout JSON export
"""

import json
import logging
from pathlib import Path
from typing import Optional

from .config import PipelineConfig
from .ocr_extractor import OCRExtractor
from .layout_analyzer import LayoutAnalyzer
from .structured_parser import StructuredParser, JudgementSchema
from .template_learner import TemplateLearner
from .section_retriever import SectionIndex, SectionRetriever
from .constrained_generator import StructuredGenerator
from .document_reconstructor import DocumentReconstructor

logger = logging.getLogger(__name__)


class JudgementReplicationPipeline:
    """
    End-to-end pipeline for Gujarati Legal Judgement Replication.

    Two modes:
    1. INDEXING: Process reference judgements to build templates and indices
    2. GENERATION: Generate new judgement replicating the learned structure

    Layout capture runs automatically during process_single_pdf if
    config.layout.save_visual_overlay / save_layout_json / save_table_json are True.
    """

    def __init__(self, config: Optional[PipelineConfig] = None):
        self.config = config or PipelineConfig()
        self.config.ensure_dirs()

        self.ocr = OCRExtractor(self.config.ocr)
        self.layout_analyzer = LayoutAnalyzer(self.config.layout)
        self.parser = StructuredParser(self.config.parser)
        self.template_learner = TemplateLearner(self.config)
        self.section_index = SectionIndex(self.config.retrieval)
        self.retriever = SectionRetriever(self.config.retrieval, self.section_index)
        self.generator = StructuredGenerator(self.config.generation)
        self.reconstructor = DocumentReconstructor(self.config)

    # =========================================================================
    # INDEXING MODE
    # =========================================================================

    def process_single_pdf(self, pdf_path: Path) -> JudgementSchema:
        """
        Process a single PDF through OCR → Layout → Structure.

        Layout capture outputs (controlled by config.layout flags):
          Option 1: PNG overlay  → data/layout_overlays/<stem>_page_XX_overlay.png
          Option 2: Table JSON   → data/tables/<stem>_tables.json
          Option 3: Layout JSON  → data/layout_json/<stem>_layout.json
        """
        pdf_path = Path(pdf_path)
        logger.info(f"Processing: {pdf_path.name}")

        # Phase 1: OCR
        logger.info("  Phase 1: OCR extraction...")
        ocr_doc = self.ocr.extract(pdf_path)
        logger.info(f"  Extracted {ocr_doc.total_pages} pages")

        # Phase 2: Layout analysis + capture
        logger.info("  Phase 2: Layout analysis + capture...")

        overlay_dir = (
            self.config.layout_overlay_dir
            if self.config.layout.save_visual_overlay
            else None
        )
        layout_json_dir = (
            self.config.layout_json_dir
            if self.config.layout.save_layout_json
            else None
        )
        table_json_dir = (
            self.config.table_json_dir
            if self.config.layout.save_table_json
            else None
        )

        layout = self.layout_analyzer.analyze(
            ocr_doc,
            pdf_path=pdf_path,
            overlay_dir=overlay_dir,
            layout_json_dir=layout_json_dir,
            table_json_dir=table_json_dir,
        )

        logger.info(
            f"  Sections: {len(layout.sections)} | "
            f"Headers: {len(layout.detected_headers)} | "
            f"Tables: {len(layout.tables)}"
        )

        # Phase 3: Structured parsing
        logger.info("  Phase 3: Structured parsing...")
        schema = self.parser.parse(layout, source_file=str(pdf_path))

        output_name = pdf_path.stem + "_parsed.json"
        schema.save(self.config.parsed_dir / output_name)
        logger.info(f"  Saved schema: {output_name}")

        return schema

    def index_reference_corpus(self, pdf_dir: Optional[Path] = None):
        """
        Process all reference PDFs and build:
        - Structured schemas
        - Templates
        - Section indices

        This is the TRAINING phase.
        If parsed JSONs already exist in data/parsed/, they are reused (no re-OCR).
        """
        schemas = []

        # Load already-parsed JSONs
        parsed_dir = self.config.parsed_dir
        existing_jsons = list(parsed_dir.glob("*_parsed.json"))
        if existing_jsons:
            logger.info(f"Loading {len(existing_jsons)} pre-parsed schemas from {parsed_dir}...")
            for json_path in existing_jsons:
                try:
                    schema = JudgementSchema.load(json_path)
                    schemas.append((schema, str(json_path)))
                    logger.info(f"  Loaded: {json_path.name}")
                except Exception as e:
                    logger.error(f"  Failed to load {json_path.name}: {e}")

        # Process any PDFs not yet parsed
        pdf_dir = Path(pdf_dir) if pdf_dir else self.config.input_dir
        pdf_files = list(pdf_dir.glob("*.pdf"))
        parsed_stems = {p.stem.replace("_parsed", "") for p in existing_jsons}

        for pdf_path in pdf_files:
            if pdf_path.stem in parsed_stems:
                continue
            try:
                schema = self.process_single_pdf(pdf_path)
                schemas.append((schema, str(pdf_path)))
            except Exception as e:
                logger.error(f"Failed to process {pdf_path.name}: {e}")
                continue

        if not schemas:
            logger.error("No documents were successfully processed or loaded")
            return

        # Phase 4: Template Learning
        logger.info("Phase 4: Learning templates...")
        for schema, _ in schemas:
            self.template_learner.learn_from_schema(schema)
        self.template_learner._build_templates()
        self.template_learner.save_templates(self.config.template_dir)
        logger.info(f"  Learned {len(self.template_learner.templates)} templates")

        # Phase 5: Build Section Index
        logger.info("Phase 5: Building section indices...")
        self.section_index.build_index(schemas)
        index_path = Path(self.config.retrieval.index_path)
        self.section_index.save_index(index_path)
        logger.info("  Section indices saved")

        logger.info("=" * 50)
        logger.info("INDEXING COMPLETE")
        logger.info(f"  Documents processed : {len(schemas)}")
        logger.info(f"  Templates           : {self.config.template_dir}")
        logger.info(f"  Indices             : {self.config.retrieval.index_path}")
        logger.info(f"  Layout overlays     : {self.config.layout_overlay_dir}")
        logger.info(f"  Layout JSON maps    : {self.config.layout_json_dir}")
        logger.info(f"  Table extractions   : {self.config.table_json_dir}")
        logger.info("=" * 50)

    # =========================================================================
    # GENERATION MODE
    # =========================================================================

    def generate_judgement(
        self,
        case_context: dict,
        document_type: str = "general",
        output_format: str = "docx",
    ) -> Path:
        """
        Generate a new judgement replicating the learned structure.
        """
        logger.info(f"Generating judgement for: {case_context.get('case_number', 'unknown')}")

        self._ensure_loaded()

        if "reference_text" not in case_context:
            parsed_dir = self.config.parsed_dir
            if parsed_dir.exists():
                for f in sorted(parsed_dir.glob("*_raw.txt")):
                    ref_text = f.read_text(encoding="utf-8")
                    import re
                    ref_text = re.sub(r'\n--- PAGE BREAK ---\n', '\n', ref_text)
                    ref_text = re.sub(r'Page \d+ of \d+\n?', '', ref_text)
                    case_context["reference_text"] = ref_text.strip()
                    logger.info(f"Using reference: {f.name}")
                    break

        template = self.template_learner.get_template(document_type)
        if not template:
            from .template_learner import DocumentTemplate
            template = DocumentTemplate(
                document_type=document_type,
                section_order=["general"],
            )

        logger.info("Phase 6: Generating by format replication...")
        schema = self.generator.generate_full_judgement(
            document_template=template,
            case_context=case_context,
            retriever=self.retriever,
        )

        logger.info("Phase 7: Saving output...")
        case_num = case_context.get("case_number", "output")
        safe_name = case_num.replace("/", "_").replace("\\", "_")
        full_text = "\n".join(s.content for s in schema.sections)

        if output_format == "text":
            output_path = self.config.output_dir / f"judgement_{safe_name}.txt"
            output_path.parent.mkdir(parents=True, exist_ok=True)
            output_path.write_text(full_text, encoding="utf-8")
        elif output_format == "docx":
            output_path = self.config.output_dir / f"judgement_{safe_name}.docx"
            output_path.parent.mkdir(parents=True, exist_ok=True)
            self._write_plain_docx(full_text, output_path)
        elif output_format == "json":
            output_path = self.config.output_dir / f"judgement_{safe_name}.json"
            output_path.parent.mkdir(parents=True, exist_ok=True)
            schema.save(output_path)
        else:
            raise ValueError(f"Unsupported output format: {output_format}")

        logger.info(f"Generated judgement saved to: {output_path}")
        return output_path

    def _write_plain_docx(self, text: str, output_path: Path):
        """Write generated text to DOCX with basic Gujarati formatting."""
        from docx import Document
        from docx.shared import Pt, Inches
        import re

        doc = Document()
        section = doc.sections[0]
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.5)
        section.right_margin = Inches(1.0)

        for line in text.split("\n"):
            stripped = line.strip()
            if not stripped:
                continue
            para = doc.add_paragraph()
            parts = re.split(r'\*\*(.+?)\*\*', stripped)
            for j, part in enumerate(parts):
                if not part:
                    continue
                run = para.add_run(part)
                run.font.name = "Shruti"
                run.font.size = Pt(12)
                if j % 2 == 1:
                    run.bold = True

        doc.save(str(output_path))

    def _ensure_loaded(self):
        """Ensure templates and indices are loaded."""
        if not self.template_learner.templates:
            template_dir = self.config.template_dir
            if template_dir.exists() and list(template_dir.glob("template_*.json")):
                self.template_learner.load_templates(template_dir)
            else:
                logger.warning("No templates found. Run index_reference_corpus first.")

        if not self.section_index._initialized:
            index_path = Path(self.config.retrieval.index_path)
            if index_path.exists() and (index_path / "metadata.json").exists():
                self.section_index.load_index(index_path)
            else:
                logger.warning("No section index found. Run index_reference_corpus first.")

    # =========================================================================
    # REPLICATION MODE
    # =========================================================================

    def replicate_structure(
        self,
        source_pdf: Path,
        new_case_context: dict,
        output_format: str = "docx",
    ) -> Path:
        """
        Replicate the exact structure of a source judgement with new case content.
        """
        logger.info(f"Replicating structure from: {source_pdf}")

        source_schema = self.process_single_pdf(source_pdf)

        from .template_learner import DocumentTemplate, SectionTemplate

        source_template = DocumentTemplate(
            document_type="replicated",
            section_order=[s.section_type for s in source_schema.sections],
        )

        for section in source_schema.sections:
            section_template = SectionTemplate(
                section_type=section.section_type,
                typical_titles=[section.title],
                avg_length=len(section.content),
                numbering_style=section.numbering or "",
                typical_subsection_count=len(section.subsections),
                example_openings=[section.content[:100]],
                example_closings=[section.content[-100:]],
            )
            source_template.section_templates[section.section_type] = section_template

        self._ensure_loaded()

        schema = self.generator.generate_full_judgement(
            document_template=source_template,
            case_context=new_case_context,
            retriever=self.retriever,
        )

        case_num = new_case_context.get("case_number", "replicated")

        if output_format == "docx":
            output_path = self.config.output_dir / f"judgement_{case_num}.docx"
            self.reconstructor.reconstruct_to_docx(schema, output_path)
        elif output_format == "text":
            output_path = self.config.output_dir / f"judgement_{case_num}.txt"
            text = self.reconstructor.reconstruct_to_text(schema)
            output_path.write_text(text, encoding="utf-8")
        else:
            output_path = self.config.output_dir / f"judgement_{case_num}.json"
            schema.save(output_path)

        logger.info(f"Replicated judgement saved to: {output_path}")
        return output_path